from pathlib import Path
import datetime as dt
import hashlib
import json
import re
import subprocess
import zipfile
from docx import Document

base=Path(__file__).resolve().parents[1]
root=base.parents[1]
evidence=base/'evidence'
doc=base/'Dragon_Market_Reviewer_Guide_and_Signoff_Form_COMPLETED_CODEX.docx'
d=Document(doc)
errors=[]
activity_count=0
for i in [11,12,16,17,20,22]:
    for row in d.tables[i].rows[1:]:
        activity_count+=1
        if row.cells[1].text.count('☒')!=1:errors.append('Invalid status '+row.cells[0].text[:12])
for i,t in enumerate(d.tables):
    for j,r in enumerate(t.rows):
        for k,c in enumerate(r.cells):
            if re.search(r'_{4,}|Click here to type',c.text):errors.append(f'Unfilled table {i} row {j} cell {k}')
assert activity_count==68
assert any(p.text.startswith('Reviewer signature: /s/ Codex') for p in d.paragraphs)
assert any(p.text.startswith('☒ REJECT / CHANGES REQUIRED') for p in d.paragraphs)
before=json.loads((evidence/'source-hashes-before.json').read_text())
after={p:hashlib.sha256((root/p).read_bytes()).hexdigest() for p in before}
changed=[p for p in before if before[p]!=after[p]]
errors.extend('Source modified: '+p for p in changed)
assert not errors,errors
cleanup=json.loads((evidence/'cleanup.json').read_text())
assert cleanup['running_containers_after_review']==[]
browser=json.loads((evidence/'browser/observations.json').read_text())
assert browser['completed'] is True
assert not browser['page_errors']
results=json.loads((evidence/'check-results.json').read_text())
assert len(results['checks'])==68
validation={'utc':dt.datetime.now(dt.timezone.utc).isoformat(),'docx_opens':True,'activity_cells_with_one_status':activity_count,'unfilled_required_fields':0,'researcher_signature_not_forged':True,'all_original_tracked_files_unchanged':not changed,'tracked_files_checked':len(before),'source_sha':subprocess.check_output(['git','rev-parse','HEAD'],cwd=root,text=True).strip(),'browser_completed':browser['completed'],'application_changes':'none','errors':errors}
(evidence/'deliverable-validation.json').write_text(json.dumps(validation,indent=2),encoding='utf8')
files=sorted(p for p in base.rglob('*') if p.is_file() and not any(x in p.relative_to(base).parts for x in ['browser-deps','__pycache__']) and p.name!='SHA256SUMS.txt')
manifest=''.join(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(base).as_posix()+'\n' for p in files)
(base/'SHA256SUMS.txt').write_text(manifest,encoding='utf8')
files.append(base/'SHA256SUMS.txt')
archive=root/'Dragon_Market_Codex_Review_Evidence_2026-10-21.zip'
with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in files:z.write(p,'Dragon_Market_Codex_Review_2026-10-21/'+p.relative_to(base).as_posix())
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    for line in manifest.splitlines():
        digest,name=line.split('  ',1)
        assert hashlib.sha256(z.read('Dragon_Market_Codex_Review_2026-10-21/'+name)).hexdigest()==digest,name
archive_digest=hashlib.sha256(archive.read_bytes()).hexdigest()
(root/(archive.name+'.sha256')).write_text(archive_digest+'  '+archive.name+'\n',encoding='utf8')
print(json.dumps({'archive':str(archive),'files':len(files),'bytes':archive.stat().st_size,'sha256':archive_digest,'validation':validation},indent=2))
