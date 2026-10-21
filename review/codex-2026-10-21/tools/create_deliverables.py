from pathlib import Path
from collections import Counter
import datetime as dt
import hashlib
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
E=BASE/'evidence'
ROOT=BASE.parents[1]
SHA='f4ba21185d1ba7b4bbb56f5addcc27d1b75bdec7'
NOW=dt.datetime.now(dt.timezone.utc)
SIGNED=NOW.isoformat(timespec='seconds')
DATE=NOW.date().isoformat()
IDENTITY=json.loads((E/'identity.json').read_text())
BROWSER=json.loads((E/'browser/observations.json').read_text())
PURCHASE=BROWSER['purchase']
SIGNATURE='/s/ Codex (OpenAI AI coding assistant)'

index={
'E01':'identity.json; source-hashes-before.json; source-integrity.json; source-hashes-after.json',
'E02':'executions.json; platform.json; *-version.log; go-cgo.log; gcc-version.log',
'E03':'go-build.log; go-vet.log; go-tests.log; go-race.log',
'E04':'full-verifier.log; full-verifier/results.json; full-verifier/frontend-*.log',
'E05':'wallet-regressions.log; baseline-verifier-retry.log; baseline-verifier-retry/verification.json',
'E06':'purchase-mutations.log; purchase-mutations/; wallet-mutations.log; wallet-mutations/',
'E07':'playwright.log; playwright-results.json; playwright-test-results/; playwright-report/',
'E08':'browser/observations.json; browser/purchase-example.json; browser/browser-trace.zip; browser/*.png',
'E09':'independent-tests.log; independent-tests.json; ../tools/independent_postgres_test.go.txt',
'E10':'browser/axe-wallet-desktop.json; browser/axe-wallet-mobile.json; browser/wallet-mobile.png',
'E11':'source-excerpts.txt; review-notes.md',
'E12':'reviewer-incident.json; baseline-verifier/; baseline-verifier.log; browser-first-attempt/',
'E13':'independent-tests-oracle.log; independent-tests-oracle.json',
'E14':'historical-evidence-restoration.json; source-integrity.json; docker-before.log; cleanup.json',
}

checks={}
def add(prefix, rows):
    for n,(status,evidence,note) in enumerate(rows,1):
        checks[f'{prefix}-{n:02}']={'status':status,'evidence':evidence,'observation':note}
add('AUTO',[
('P','E01','Exact target SHA; initially clean existing checkout. No clone/pull; requester-authorized deviation.'),
('P','E03','Original Go build and vet exit 0.'),
('P','E03','Uncached verbose suite exit 0; real PostgreSQL Testcontainers; no skipped tests observed.'),
('P','E03','Windows CGO race suite exit 0; GCC 16.1.0 available.'),
('P','E04','Full verifier exit 0: build/vet/tests/race plus frontend typecheck/lint/15 tests/build.'),
('P','E05, E12','Focused daily-limit tests and final baseline verifier pass. First verifier attempt failed due to reviewer file packaging; preserved and corrected.'),
('P','E06','All three daily-limit mutants rejected by named SQL-backed assertion; two purchase mutants also detected.'),
('P','E07','All 7 existing browser tests passed with installed Chrome and isolated disposable DB.'),
])
add('UI',[
('P','E08','System Status shows development, exact SHA, PostgreSQL identity, successful read and HTTP health.'),
('P','E08','Switched all three seeded guilds; initial visible 100.00 G; explicitly not authentication.'),
('P','E08','Created uniquely named rare item; owner/type/700-unit price verified in real snapshot.'),
('P','E08','Listed for 700; zero price blocked by native range validation; invalid API price rejected.'),
('P','E08','Buyer 10000→9300, seller 10000→10700; owner transferred once; sold rebuy rejected.'),
('P','E08','Common-item listing cancelled to free; sold listing cancellation rejected (400).'),
('P','E08','Legendary 700-unit base item; active auction, backend deadline approximately 24 hours ahead.'),
('P','E08','500 bid reserves 500; total gold unchanged; available gold decreases by hold.'),
('P','E08','Raise 500→525 uses delta 25; bidder 552 outbids; old hold released; highest cancellation succeeds; repeated cancellation 409.'),
('P','E08','Same-key replay: 204 plus replay header; unchanged ownership/balances; exactly 2 audit rows and 1 idempotency record.'),
('P','E08, E09, E11','Inspected raw persisted pre/post values and date; distinguished them from GET-wallet projection. DM-REV-CODEX-001 documents wrong Wallet display.'),
('P','E08','526 rejected below required 552 with visible error; zero base price HTTP 400; no false success.'),
('F','E10; DM-REV-CODEX-003','Mobile audit scroller lacks keyboard focus (axe serious). Desktop scan, skip link, route refresh and page-width checks passed.'),
('P','E07, E08, E11','Local harness executes C/F/G in isolated DBs; hosted gateway tests reject harness/arbitrary debug routes. No live hosted site assessed.'),
])
add('WAL',[
('P','E05, E11','Supplied S(current UTC day)+all holds≤positive L is coherent under valid starting state, fixed limits and reliable UTC clock; L=0 unlimited. Human intent confirmation remains owner-owned.'),
('P','E03, E05','Funds/capacity checks and row locks enforce reserve allowance; failed operations have no committed effect.'),
('P','E05, E06','Old hold release reduces reserved only; mixed-age test and mutation rejection preserve completed spending.'),
('P','E05','500 hold at 500 limit rejects 300 spend before release; Tehran midnight does not reset UTC allowance.'),
('P','E05','After release, 300 debit persists; next 500 rejected; SQL gold=9700/reserved=0/completed=300/audits=2.'),
('P','E05','Old settlement exchanges reserved for spending; SQL gold debited once; repeated settlement rejected.'),
('P','E05, E08','Cancellation/outbid restores reservation capacity while completed 700-unit purchase remains charged.'),
('P','E09','Independent SQL test: spend 500, earn 1000, next spend 1 rejected; gold=10500/spent=500/audits=2.'),
('P','E03, E05','Exact boundary accepted; over-cap rejected; zero-limit behavior and integer limits exercised.'),
('P','E03, E05','Positive amounts, insufficient funds/reserve and int64 overflow checked in domain suites.'),
('P','E05, E09, E11','GET wallet projects prior-day counter to 0 without write; raw inspector retains 300 and prior date. UI defect separately reported.'),
('P','E08, E09','Independently queried PostgreSQL wallet/item/audit rows and debit totals; no reliance on UI responses alone.'),
])
add('TX',[
('P','E03, E07','Concurrent-buy and Research Lab tests: single winner, intended owner and committed ledger effects.'),
('P','E03, E07','Controlled late failure after debit/earn/transfer rolls back every persisted effect and idempotency claim.'),
('P','E03, E08','Buy/bid replay suites inspect SQL; UI exact purchase replay produces no extra ledger entries.'),
('P','E03, E08','New-key attempt on sold order rejected; changed same-key payload covered by existing SQL-backed test.'),
('P','E03','Concurrent auction starts yield one active auction; partial unique index and item transition reviewed.'),
('P','E03, E08','ceil(current/20) minimum increment reviewed; 525→552 boundary and 526 rejection observed; overflow test passes.'),
('P','E03, E08','Outbid release and new hold within outer transaction; concurrent-bid SQL test leaves one active highest reservation.'),
('P','E05, E08','Same-bidder delta reserve confirmed: total reservation 525, not 1025.'),
('P','E03, E08','Only eligible active highest cancellation releases correct wallet; repeated invalid cancellation 409.'),
('P','E03, E11','Controlled-clock auction test verifies final-five-minute extension; default config is 5-minute window/extension.'),
('P','E03, E05, E07','Worker plus concurrent settlement driver and repeat batch produce one debit/credit/ownership transfer.'),
('P','E05','12 competing reserve/spend operations accept exactly 5×100 under limit 500; SQL usage=500.'),
('P','E03, E05','Injected audit failure rolls back accounting and stored day; purchase late-failure test rolls back ownership.'),
('F','E09, E11; DM-REV-CODEX-002','Nested transactions reuse outer transaction, but buyer-then-seller wallet locks deadlock reciprocal purchases (40P01); one rolls back. No automatic retry.'),
])
add('DB',[
('P','E03, E07, E08','Fresh Testcontainers/E2E apply migrations 1–11 before seeding; v2 wallets and successful API startup observed.'),
('P','E03, E05','Legacy initialization and focused migration tests reject populated v1 without changed balances/counters/ledger.'),
('P','E09','Reviewer ran down migration on populated v2; explicit refusal; marker=2/gold=9877/spent=123 preserved.'),
('P','E09','Reviewer temp v1 schema: new wallet read rejects absent daily_accounting_version (42703); no silent fallback.'),
('P','E01, E14','All 234 tracked files byte-identical at completion, including historical FAIL/PASS evidence; temporary verifier output separately captured and restored.'),
('P','E14','Only positively identified disposable containers removed; no Compose launch/reset/down -v; zero running containers before and after.'),
('P','E03, E05','Tehran midnight vs UTC midnight and persisted UTC day covered; carried reservations retained.'),
('P','E13','Reviewer real HTTP oracle cases valid/zero/negative/malformed/503/timeout preserve cached 250; SQL base=250, fixed listing stays 100.'),
('P','E03, E11','Go build/test commands run independently of frontend tooling. Frontend formats/validates/suggests bids; server remains monetary authority.'),
('P','E01, E02, E08','OS/tool/browser/PostgreSQL versions, SHA, timestamps, argv, exits, source hashes and deviations recorded.'),
('N/A','E01, E08','Hosted deployment not used. Local stack independently reports exact target SHA; no claims about Render/Neon/live CI.'),
('P','E03, E11','Local-only dev snapshot and harness restrictions plus gateway gate/origin/path tests reviewed. No production security certification.'),
])
for i in range(1,9):
    checks[f'TASK-{i:02}']={'status':'NR','evidence':'E11; scope exclusion','observation':'Separate research repository, exact task packages, reference patches and controlled hidden oracles not supplied/reviewed. No task/oracle/fairness/leakage approval claimed.'}

findings=[
{'id':'DM-REV-CODEX-001','severity':'Medium','tests':'UI-11 / WAL-11 display limitation (backend projection itself passes)',
 'title':'Wallet shows yesterday’s spending as “Spent today”',
 'reproduction':'On the positively labelled disposable E2E DB only, set guild-buyer daily_spent=300 and daily_spent_on to UTC yesterday. Open /#wallet, select guild-buyer and reload. GET /api/guilds/guild-buyer/wallet independently.',
 'expected':'Today’s visible completed spend is 0.00 G; raw inspector retains yesterday’s 300 with its date. Reading must not alter stored history.',
 'actual':'Wallet visibly shows Spent today 3.00 G. GET wallet returns daily_spent=0; SQL retains 300 and yesterday’s date. Reproduced in both custom browser sessions and reviewer SQL test. No spending-limit bypass observed.',
 'source':'web/src/components.tsx:83; web/src/App.tsx:133 and :919; internal/postgres/inspection.go:42; internal/guild/service.go:GetGuild',
 'evidence':'E08 browser/finding-stale-spent-today.png, browser/observations.json (stale_counter); E09 independent-tests.log',
 'action':'Use current-day wallet projection for Wallet/Overview, or project with the observation’s UTC day while keeping dated raw inspector data. Add a stale-day UI regression.',
 'status':'Open; exact source unchanged.'},
{'id':'DM-REV-CODEX-002','severity':'Medium','tests':'TX-14',
 'title':'Reciprocal purchases acquire wallet locks in opposite order and deadlock',
 'reproduction':'Run tools/run_independent.py against this exact checkout. The retained source copy adds only reviewer_independent_test.go. Create a rare listing owned by seller and another owned by buyer. Launch buyer buying seller’s item and seller buying buyer’s item concurrently. A test-only wallet wrapper waits after each real Spend has locked its buyer row; then lets both real Earn calls proceed. This deterministically schedules a naturally possible interleaving without changing application SQL or domain code.',
 'expected':'Valid reciprocal operations complete through consistent multi-wallet lock order or a bounded transaction retry; failed attempts must roll back.',
 'actual':'TestReviewerReciprocalPurchasesDoNotDeadlock fails: one service operation returns ERROR: deadlock detected (SQLSTATE 40P01), the other succeeds. SQL total gold=30000, audit rows=2, sold orders=1. Rollback/money conservation hold. HTTP 500 follows from handler mapping by source inspection; HTTP deadlock response was not separately executed.',
 'source':'internal/order/service.go:132, :135; internal/postgres/guild_repo.go:60; internal/postgres/tx.go:37; internal/api/order_handler.go:188',
 'evidence':'E09 independent-tests.log; tools/independent_postgres_test.go.txt; E11 source-excerpts.txt',
 'action':'Acquire all participating wallet locks in a deterministic order, or deliberately retry the entire transaction on 40P01/40001 with bounded attempts. Retain rollback/idempotency guarantees. Review cross-auction two-wallet paths too.',
 'status':'Open; deterministic counterexample preserved; no production source patched.'},
{'id':'DM-REV-CODEX-003','severity':'Medium','tests':'UI-13',
 'title':'Mobile wallet audit region cannot receive keyboard focus',
 'reproduction':'With wallet audit rows present, open /#wallet at 390×844 in Chrome. Run axe-core 4.11.1 WCAG2 A/AA and WCAG2.1 AA scan; inspect .table-wrap.',
 'expected':'A horizontally scrollable audit table has an accessible keyboard scrolling path, with suitable region naming/focus behavior.',
 'actual':'Repeated scans flag scrollable-region-focusable, impact serious, target .table-wrap. Audit table has overflow content with no focusable descendant/tabindex. Desktop scan has zero violations. Skip link, main focus, mobile page width, navigation and refresh checks pass. No screen-reader evaluation or complete WCAG certification performed.',
 'source':'web/src/App.tsx:1181 (AuditTable wrapper); web/src/styles.css (.table-wrap)',
 'evidence':'E10 browser/axe-wallet-mobile.json and browser/axe-wallet-desktop.json; browser/wallet-mobile.png; E08 browser/browser-trace.zip',
 'action':'Provide a named keyboard-focusable scroll region and verify horizontal scrolling and focus visibility with keyboard and assistive technology.',
 'status':'Open; reproduced twice; no fix applied.'},
]

summaries={p:Counter(v['status'] for k,v in checks.items() if k.startswith(p+'-')) for p in ['AUTO','UI','WAL','TX','DB','TASK']}
notes='''# Reviewer source reasoning and limits

Reviewer: Codex, an OpenAI AI coding assistant. This is a new AI-executed review, not an independent human review, institutional attestation or authority to freeze a research baseline. No source changes were made. Human independence from earlier AI-assisted authorship is not established.

The supplied positive-limit policy S(d)+R<=L is internally coherent. Reserve/spend check remaining capacity using guarded subtraction; nonnegative SQL constraints and funds checks bound additions. Release decreases only R. Deduct exchanges a hold for equal completed spending while decreasing gold. Earn changes gold only. UTC rollover drops nonnegative completed spending and preserves all holds. Thus each transition preserves the bound from valid service-created state with a fixed limit and reliable clock. L=0 still has int64 accounting bounds. This assessment accepts the supplied technical policy; it does not establish what a human researcher intends.

Wallet reads copy the row and project the day; diagnostic snapshots intentionally return raw values and dates. The UI incorrectly labels the raw historical counter as today's spending, independently reproduced. The inspector's visible comparison also omits the raw date; readers must use the recorded snapshot/date instead of assuming a current-day projection. The audit column says UTC time but prints offset-bearing raw timestamps (for example +03:30); this is an additional low-severity labeling observation, not a debit/accounting error.

Wallet UPDATE and audit insert share the context transaction. Order operations lock item then order; same-item competition serializes. Auction operations lock the auction and reuse that outer transaction. The wallet repository always uses FOR UPDATE; nested RunInTransaction calls reuse the transaction. There is no ordering across multiple distinct wallet rows and no 40P01 retry. Opposing purchases can deadlock despite same-listing contention tests passing. The reviewer barrier only schedules after real successful Spend calls and does not manufacture a lock type absent from production. The observed abort preserves atomicity; this is a liveness/error-handling finding rather than evidence of money duplication.

The price increment is current + ceil(current/20), with overflow rejected. The browser exercised a 525->552 minimum and a 526 rejection; existing controlled-clock tests cover extension and repeated/concurrent settlement. Reviewer oracle checks use a real local HTTP server and the real resilient oracle, then persist its result through the real repository to compare base and fixed-order prices. They do not claim a separate new updater-scheduling test; the existing updater test was also executed.

Migration 11's refusal is an intentional fresh-data boundary. The reviewer executed populated downgrade refusal and old-schema rejection in isolated tests. Existing legacy initialization checks compare surviving balances and ledger; no populated database was cleared or reinterpreted.

Test adequacy limits: ordinary passing suites did not cover reciprocal wallet-lock ordering, dated-wallet rendering or mobile audit keyboard scrolling. Broader cross-auction/cross-order interleavings, worker batch starvation after persistent failure, long transactions spanning UTC midnight, clock skew, a complete screen-reader audit, other browsers/OSs and live cloud deployment/CI were not independently exercised. Simulation uses injected UTC clocks; the review did not wait for real midnight. No load/performance or formal correctness proof is claimed.

Section 8: the separate research repository, exact T001/T002 packages, human references and controlled private evaluation oracles were not supplied in the reviewed application checkout and were not accessed. All eight task checks are NOT RUN. No statement about fairness, feature novelty, hidden-oracle adequacy or absence of leakage is made. No research pin, branch, tag, hosted database or deployment changed.

Environment: existing initially clean exact-SHA checkout as explicitly requested, not a newly cloned independent machine; Windows 11 build 26200, Go 1.26.0 windows/amd64, Node 24.18.0/npm 11.16.0, GCC 16.1.0, Docker Desktop 4.58.0/Engine 29.1.5/Compose 5.0.1, PostgreSQL 16.14 Linux container, Chrome 155.0.8059.85. Browser channel chrome selected instead of downloading Chromium, as README permits. Workbench frontend/backend exercised through the repository's disposable E2E stack at 5175/8085/8087; shared Compose stack at 5173 was not started. Testcontainers used separate fresh databases. No remote CI status was queried or inherited from old logs.

Reviewer setup installed python-docx/reportlab/pypdf for document work and isolated @axe-core/playwright 4.11.1 for accessibility scanning. npm ci installed the locked frontend dependencies; no lockfile change; npm audit reported zero vulnerabilities at execution (not a security certification). npm emitted an unknown user node-linker config warning, jsdom reported unimplemented scrollTo, and Playwright emitted color-environment warnings; these did not fail the relevant checks. Early document-text extraction failed under cp1252 and was rerun with Python UTF-8. Some read-only exploratory file searches used nonexistent paths or Windows wildcard forms and were corrected; none changed application state.

The first baseline-verifier invocation was contaminated by the review author's temporary .go file under review/tools, outside its intended postgres package. Build/vet then reported undefined helper symbols in that review-only package. This was a reviewer packaging mistake, not an application regression. The file was renamed .go.txt and only copied into a disposable source tree for the additional tests. A complete fresh baseline-verifier retry passed; both runs are retained. Original tracked historical evidence was byte-backed-up before each exact verifier invocation, copied separately for this review, then restored; source hashes confirm all originals unchanged.

The first custom browser script completed purchase/replay/auction/accessibility/stale-day observations but its final exact Research Lab locator omitted the DEV badge in the accessible name and timed out. This was a reviewer locator mistake. The entire first attempt was preserved; a complete second fresh-disposable-stack run used the documented #lab route and completed. Both browser sessions reproduced the same two UI findings; the concurrency counterexample was independently reproduced twice in SQL-backed Go tests. Existing seven-test Playwright suite passed independently before these custom sessions.

The combined final reviewer-authored SQL-backed suite executed all six additional checks: five passed and the reciprocal-purchase expectation failed again with 40P01. The initial five-test run and supplemental oracle run remain preserved. Optional Word COM PDF export did not complete; a complete reflowed PDF was instead rendered from the filled DOCX with ReportLab, with the rendering method disclosed on its first page. The stalled optional pymupdf download was stopped; pypdfium2 was installed for visual inspection. No application check result was changed by document rendering.

No implementation fixes were made, no messages/issues were posted externally, no repository was pulled/cloned, and no baseline owner decision was supplied. Review-scoped typed signature only; no certificate-backed or human signature is claimed.
'''
(E/'review-notes.md').write_text(notes,encoding='utf8')
(E/'check-results.json').write_text(json.dumps({'sha':SHA,'reviewer':SIGNATURE,'signed_at_utc':SIGNED,'checks':checks,'summary':{k:dict(v) for k,v in summaries.items()},'findings':findings},indent=2,ensure_ascii=False),encoding='utf8')

report=[f'# Dragon Market — completed AI reviewer report\n\nDecision: **REJECT / CHANGES REQUIRED within reviewed application scope**.\n\nReviewed commit: `{SHA}`.\n\nReviewer: **Codex (OpenAI AI coding assistant)**. Typed acknowledgement: `{SIGNATURE}` at {SIGNED}. This signature identifies AI authorship; no human, OpenAI corporate or baseline-owner sign-off is represented.\n\nThe existing automated Go, race, frontend, mutation and seven-test browser suites passed. Additional reviewer checks exposed three unresolved issues. Money conservation and rollback held in the deadlock test. No financial-invariant bypass was observed. The application may remain useful as a local demo; this review does not recommend freezing it as a fully approved research baseline until the findings are resolved or explicitly dispositioned by the owner.\n',
'## Review scope and actual execution\n\n'+notes.split('Environment:')[1].split('Reviewer setup')[0].strip()+'\n',
'## Findings\n']
for f in findings:
    report.append(f"### {f['id']} — {f['title']} ({f['severity']})\n\n"+'\n\n'.join(f'**{k.capitalize()}:** {f[k]}' for k in ['tests','reproduction','expected','actual','source','evidence','action','status'])+'\n')
report.append('## Section counts\n\n| Section | Total | P | F | B | NR | N/A |\n|---|---:|---:|---:|---:|---:|---:|')
for prefix,c in summaries.items():report.append(f"| {prefix} | {sum(c.values())} | {c['P']} | {c['F']} | {c['B']} | {c['NR']} | {c['N/A']} |")
report.append('\nThe failed UI rendering case is also recorded as DM-REV-CODEX-001; WAL-11 passes for the backend projection itself. The row counts above count form activities, not every subassertion or finding.\n\n## Complete activity record\n')
for k,v in checks.items():report.append(f"- **{k}: {v['status']}** — {v['observation']} Evidence: {v['evidence']}.")
report.append('\n## Evidence index\n\nAll paths below are relative to `evidence/` in this package. Every artifact is covered by SHA256SUMS.txt.\n')
for k,v in index.items():report.append(f'- **{k}**: {v}')
report.append('\n## Full source reasoning, exceptions and chronology\n\n'+notes.split('\n\n',1)[1])
report.append(f'\n## Signature and reserved owner fields\n\n{SIGNATURE}\n\nUTC: {SIGNED}\n\nI acknowledge only the AI-executed observations, explicit code reasoning and exclusions recorded here. Final human approval remains pending; no approved SHA, pin update or owner signature is asserted. No private contact address is available for this AI reviewer.\n')
(BASE/'REVIEW_REPORT.md').write_text('\n'.join(report),encoding='utf8')

d=Document(BASE/'Guide.original.docx')
d.core_properties.author='Codex (OpenAI AI coding assistant)'
d.core_properties.last_modified_by='Codex (OpenAI AI coding assistant)'
d.core_properties.title='Dragon Market — Completed AI Review and Typed Sign-off'
d.core_properties.subject=f'Actual executed review of {SHA}; changes required; human approval pending'
d.core_properties.comments='AI reviewer-authored record. Original guide preserved separately. No researcher acceptance signed.'

def cell_text(cell,text):
    cell.text=text
    for p in cell.paragraphs:
        for run in p.runs:run.font.size=Pt(9)
        p.paragraph_format.space_after=Pt(3)
def fill_table(i,values):
    for row,value in zip(d.tables[i].rows,values):cell_text(row.cells[1],value)

d.tables[0].rows[2].cells[1].text='AI REVIEW: CHANGES REQUIRED. Independent human approval remains pending.'
fill_table(2,[
'Codex (OpenAI AI coding assistant)',
'AI reviewer in this session; no personal contact address or institutional appointment. Does not sign on behalf of OpenAI.',
DATE+' UTC',
'☒ UI / functional   ☒ Go/backend   ☒ PostgreSQL   ☒ Reproducibility   ☒ Accessibility (limited automated/keyboard scope)',
'☐ Local clean checkout   ☒ Existing local install   ☐ Hosted review URL   ☐ Source-only. Exact SHA initially clean; disposable DB harness.',
'☐ Yes   ☒ No / not established: AI reviewer requested by owner; not an independent human. No implementation authored in this review; earlier authorship independence cannot be certified.',
])
for ti in [11,12,16,17,20,22]:
    for row in d.tables[ti].rows[1:]:
        key=re.match(r'(AUTO|UI|WAL|TX|DB|TASK)-\d+',row.cells[0].text).group()
        v=checks[key]
        statuses='   '.join(('☒ ' if s==v['status'] else '☐ ')+s for s in ['P','F','B','NR','N/A'])
        cell_text(row.cells[1],statuses+'\nEvidence: '+v['evidence']+'\n'+v['observation'])
fill_table(13,[
f"Item: {PURCHASE['item_id']}\nOrder: {PURCHASE['order_id']}\nCodex-R1-2110-Emberglass; raw price 700 = 7.00 G.",
'guild-seller → guild-buyer; isolated dragon_research_disposable database.',
'Before purchase: seller gold=10000/reserved=0/spent=0; buyer gold=10000/reserved=0/spent=0; available equals gold. Limit=5000 each.',
'After purchase and same-key replay: seller gold=10700/reserved=0/spent=0; buyer gold=9300/reserved=0/spent=700; audit rows=2; idempotency rows=1. All amounts raw units.',
'guild-seller → guild-buyer once; order sold; same-key replay retains buyer; new-key rebuy rejected.',
'204 and Idempotency-Replayed:true; no new monetary, ownership or ledger effects (independent SQL).',
'E08: browser/purchase-example.json; browser/observations.json; browser/purchase-replay-inspector.png; browser/browser-trace.zip.',
])
cell_text(d.tables[18].cell(0,0),'Transaction boundary / locking issues observed\nDM-REV-CODEX-002: nested transaction reuse and rollback work, but opposite buyer/seller lock order causes 40P01 under reciprocal purchases. E09 independent-tests.log and E11 source-excerpts.txt. No financial conservation failure observed.')
cell_text(d.tables[19].cell(0,0),'Invariant counterexamples or missing concurrency interleavings\nOne naturally possible reciprocal-purchase schedule aborts one valid transaction. Existing same-item competition tests do not cover distinct items sharing two wallets in opposite directions. Cross-auction reciprocal wallets and worker starvation remain untested concerns, not claimed reproductions.')
fill_table(23,[SHA,'☒ Clean initially   ☐ Not clean. After review: additive review/ only; all 234 tracked hashes unchanged.',
'Windows 11 10.0.26200, amd64; Linux PostgreSQL via Docker Desktop.',
'Go 1.26.0 windows/amd64; Node v24.18.0; npm 11.16.0; GCC 16.1.0; CGO race passed.',
'PostgreSQL 16.14 (postgres:16-alpine); Docker Desktop 4.58.0, Engine 29.1.5, Compose 5.0.1; Chrome 155.0.8059.85.',
'☒ Isolated Testcontainers   ☒ Dedicated local disposable E2E DB   ☐ Hosted review. No Compose/shared DB touched.',
'http://127.0.0.1:5175 — disposable E2E workbench; API 8085, lab 8087. Actual displayed backend SHA: '+SHA+'. Shared 5173 launch intentionally replaced with isolated harness.',
IDENTITY['started_at_utc']+' → '+SIGNED+' (UTC; testing completed before document sign-off).',
'review/codex-2026-10-21/evidence/; E01–E14 index and SHA256SUMS.txt; packaged evidence ZIP.',
])
f=findings[1]
fill_table(24,[f['id'],'☐ Critical   ☐ High   ☒ Medium   ☐ Low   ☐ Observation',f['tests'],SHA,f['reproduction'],f['expected'],f['actual'],f['source'],f['evidence'],f['action'],'☒ Open   ☐ Fixed + verified   ☐ Accepted risk   ☐ Cannot reproduce'])
cell_text(d.tables[25].cell(0,0),'Detailed reproduction and SQL/state evidence\n'+f['actual']+'\nAll three findings, exact reproductions and evidence links are completed in Appendix C and REVIEW_REPORT.md. No finding has been fixed or accepted by the researcher.')
cell_text(d.tables[26].cell(0,0),'Concerns about test adequacy\nPassing original suites missed reciprocal-wallet deadlocks, stale-day Wallet rendering and mobile audit keyboard access. Additional cases reproduced these. Broader cross-auction interleavings, long transactions across UTC midnight, clock skew, screen-reader behavior, other platforms and live CI/cloud deployment remain unverified. T001/T002 references/oracles and prompt fairness were not reviewed. See E11 review-notes.md.')
for row,prefix in zip(d.tables[27].rows[1:],['AUTO','UI','WAL','TX','DB','TASK']):
    c=summaries[prefix]
    cell_text(row.cells[2],str(c['P']));cell_text(row.cells[3],str(c['F']))
    cell_text(row.cells[4],f"B={c['B']}; NR={c['NR']}; N/A={c['N/A']}")
fill_table(29,[
'Codex (OpenAI AI coding assistant)',
'AI reviewer: code reasoning, direct command execution, SQL-backed reproduction and browser automation; limited keyboard/axe accessibility checks. No human or corporate signing authority claimed.',
SHA,DATE+' UTC; acknowledgement time '+SIGNED,
'review/codex-2026-10-21/evidence/ and Dragon_Market_Codex_Review_Evidence_2026-10-21.zip; see SHA256SUMS.txt.',
'DM-REV-CODEX-001 (stale today display), -002 (reciprocal deadlock), -003 (mobile keyboard scroller). All open. No external issue posted.',
'Original build/vet/uncached/race suites, full verifier, focused daily-limit and final baseline verifier, five mutations, locked npm ci, frontend 15 tests/build/lint/typecheck, seven existing browser tests, custom browser/SQL/axe sessions, six reviewer-authored PostgreSQL checks. Exact commands and both reviewer-caused failed attempts retained in E02/E09/E12/E13.',
])
cell_text(d.tables[30].cell(0,0),'Decision: REJECT / CHANGES REQUIRED within the reviewed application scope.\nThree unresolved findings include a deterministic concurrency deadlock, incorrect daily-spend presentation and inaccessible mobile audit scrolling. Resolve or explicitly disposition them before baseline freeze; rerun affected checks against a new exact SHA if code changes. Existing financial safety checks passed, and aborted reciprocal purchase preserved atomicity. Optional research tasks/private oracles, independent human review, live hosted deployment, remote CI and other platforms are excluded. This is an AI reviewer acknowledgement only; human approval remains pending.')
# Every owner field gets an explicit reserved/pending disposition, never another person's signature.
fill_table(31,[
'PENDING — reserved for researcher; no owner approval supplied or inferred.',
'Not approved by owner in this review. Reviewed SHA recorded above; no approved SHA asserted.',
'Reviewer package E01–E14 supplied for owner consideration; no owner acceptance evidence recorded.',
'No research pin/branch/tag update performed; owner decision pending.',
'Not signed — researcher / baseline-owner field; outside AI reviewer authority.',
])

for p in d.paragraphs:
    if p.text.startswith('Reviewer signature:'):
        p.text=f'Reviewer signature: {SIGNATURE}    UTC date: {DATE}\nTyped electronic authorship acknowledgement at {SIGNED}; AI reviewer, not a human or certificate-backed signature.'
    elif p.text.startswith('☐ REJECT / CHANGES REQUIRED'):
        p.text=p.text.replace('☐','☒',1)
    elif re.fullmatch(r'_+',p.text.strip()):
        p.text='Completed in the corresponding result field and Appendix C; no additional unwritten response.'

# Add a prominent authorship/scope note without removing the supplied guide.
for p in d.paragraphs:
    if p.text.startswith('For invited reviewers'):
        p.insert_paragraph_before('COMPLETED AI REVIEW — Codex executed this review on '+DATE+'. No independent human review or researcher approval is claimed. Existing initially clean exact-SHA checkout used at requester instruction. All 68 activity results are explicitly completed; optional research task checks are NOT RUN. Three findings remain open. Reviewer declaration and typed acknowledgement are in Section 11.')
        break

d.add_page_break()
d.add_heading('Appendix C — Completed findings, evidence and execution exceptions',level=1)
d.add_paragraph('This appendix is part of the signed review. Evidence paths are relative to the evidence/ directory beside this document. REVIEW_REPORT.md contains the complete activity-by-activity record and detailed reasoning. Original guide preserved as Guide.original.docx.')
for f in findings:
    d.add_heading(f['id']+' — '+f['title'],level=2)
    for label in ['severity','tests','reproduction','expected','actual','source','evidence','action','status']:
        p=d.add_paragraph();p.add_run(label.capitalize()+': ').bold=True;p.add_run(f[label])
d.add_heading('Evidence index',level=2)
for key,paths in index.items():d.add_paragraph(key+' — '+paths)
d.add_heading('Source reasoning, deviations and actual execution history',level=2)
for paragraph in notes.split('\n\n')[1:]:d.add_paragraph(paragraph.strip())
for filename,caption,width in [
 ('browser/purchase-replay-inspector.png','Purchase replay: unchanged balance and ownership; independently checked ledger rows.',5.8),
 ('browser/finding-stale-spent-today.png','DM-REV-CODEX-001: prior-day 300 shown as Spent today 3.00 G; API projects zero.',2.6),
 ('browser/wallet-mobile.png','DM-REV-CODEX-003: mobile audit overflows inside the table region; axe identifies absent keyboard focus.',2.6),
]:
    d.add_paragraph(caption)
    d.add_picture(str(E/filename),width=Inches(width))
d.add_paragraph(SIGNATURE+'\nSigned UTC: '+SIGNED+'\nDecision and exclusions as in Section 11. Owner acceptance remains pending.')
output=BASE/'Dragon_Market_Reviewer_Guide_and_Signoff_Form_COMPLETED_CODEX.docx'
d.save(output)
(BASE/'SIGNATURE.txt').write_text(f'AI reviewer typed electronic acknowledgement\n{SIGNATURE}\nUTC: {SIGNED}\nReviewed SHA: {SHA}\nDecision: REJECT / CHANGES REQUIRED within reviewed scope.\nAuthorship of actual checks, findings and exclusions only. No human, institutional, OpenAI corporate or baseline-owner approval represented. No certificate-backed digital signature.\nThe adjacent SHA256SUMS.txt binds file contents for integrity; hashes alone do not authenticate a legal signature.\n',encoding='utf8')
print(json.dumps({'document':str(output),'activity_count':len(checks),'summary':{k:dict(v) for k,v in summaries.items()},'signed_utc':SIGNED},indent=2))
