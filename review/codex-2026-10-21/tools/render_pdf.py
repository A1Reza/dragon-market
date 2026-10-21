"""Render all completed DOCX content to a readable PDF, with explicit reflow."""
from pathlib import Path
from io import BytesIO
from xml.sax.saxutils import escape
import datetime,json
from docx import Document
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,Image,PageBreak
from pypdf import PdfReader

base=Path(__file__).resolve().parents[1]
input_path=base/'Dragon_Market_Reviewer_Guide_and_Signoff_Form_COMPLETED_CODEX.docx'
output_path=input_path.with_suffix('.pdf')
pdfmetrics.registerFont(TTFont('ReviewBody',r'C:\Windows\Fonts\segoeui.ttf'))
pdfmetrics.registerFont(TTFont('ReviewBold',r'C:\Windows\Fonts\segoeuib.ttf'))
pdfmetrics.registerFont(TTFont('ReviewSymbols',r'C:\Windows\Fonts\seguisym.ttf'))
pdfmetrics.registerFontFamily('ReviewBody',normal='ReviewBody',bold='ReviewBold',italic='ReviewBody',boldItalic='ReviewBold')
style=ParagraphStyle('body',fontName='ReviewBody',fontSize=9,leading=12,spaceAfter=6,splitLongWords=True)
cellstyle=ParagraphStyle('cell',parent=style,fontSize=8,leading=10,spaceAfter=2)
headstyle=ParagraphStyle('heading',parent=style,fontName='ReviewBold',fontSize=13,leading=17,spaceBefore=12,spaceAfter=8,keepWithNext=True)
substyle=ParagraphStyle('subheading',parent=headstyle,fontSize=10.5,leading=14)
def markup(s):
    s=escape(s).replace('\n','<br/>')
    for symbol in ['☐','☒','•']:
        s=s.replace(symbol,'<font name="ReviewSymbols">'+symbol+'</font>')
    return s
def para(text,sty=style):return Paragraph(markup(text),sty)
story=[para('PDF rendering note: this is a reflowed export of the complete filled Word form, produced by the reviewer with ReportLab. The editable DOCX is the original-layout deliverable. Word COM export did not complete; this PDF does not claim exact Word pagination.'),Spacer(1,8)]
doc=Document(input_path)
for item in doc.iter_inner_content():
    if isinstance(item,WordParagraph):
        if item.text.strip():
            name=item.style.name
            story.append(para(item.text,headstyle if name.startswith('Heading 1') or name=='Title' else substyle if name.startswith('Heading') else style))
        for image_node in item._p.xpath('.//a:blip'):
            embed=image_node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            blob=doc.part.related_parts[embed].blob
            from PIL import Image as PILImage
            w,h=PILImage.open(BytesIO(blob)).size
            width=min(480,630*w/h)
            story.append(Image(BytesIO(blob),width=width,height=width*h/w))
        if not item.text.strip() and not item._p.xpath('.//a:blip'):story.append(Spacer(1,4))
    elif isinstance(item,WordTable):
        count=len(item.columns)
        data=[[para(c.text,cellstyle) for c in row.cells] for row in item.rows]
        if count==1:widths=[500]
        elif count==2:
            checklist='REVIEW ACTIVITY' in item.cell(0,0).text
            widths=[260,240] if checklist else [125,375]
        elif count==3:widths=[180,160,160]
        else:widths=[230]+[(500-230)/(count-1)]*(count-1)
        table=LongTable(data,colWidths=widths,repeatRows=1 if 'REVIEW ACTIVITY' in item.cell(0,0).text else 0,splitByRow=1,splitInRow=1,hAlign='LEFT')
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#A9B7C5')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF0F6'))]))
        story.extend([table,Spacer(1,10)])

def footer(canvas,document):
    canvas.setFont('ReviewBody',7)
    canvas.setFillColor(colors.HexColor('#526270'))
    canvas.drawString(48,30,'Dragon Market · Codex AI review · f4ba21185d1b · human approval pending')
    canvas.drawRightString(564,30,str(document.page))
output=SimpleDocTemplate(str(output_path),pagesize=(612,792),leftMargin=50,rightMargin=50,topMargin=45,bottomMargin=47,title='Dragon Market — Completed AI Review',author='Codex (OpenAI AI coding assistant)')
output.build(story,onFirstPage=footer,onLaterPages=footer)
pdf=PdfReader(output_path)
text='\n'.join(p.extract_text() or '' for p in pdf.pages)
checks=json.loads((base/'evidence/check-results.json').read_text())['checks']
missing=[k for k in checks if k not in text]
assert not missing,missing
assert '/s/ Codex' in text
assert 'DM-REV-CODEX-002' in text
(base/'evidence/pdf-text.txt').write_text(text,encoding='utf8')
record={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'renderer':'ReportLab '+__import__('importlib.metadata').metadata.version('reportlab'),'source':input_path.name,'pdf':output_path.name,'pages':len(pdf.pages),'all_68_activity_ids_present':True,'signature_text_present':True,'layout':'Complete DOCX content reflowed for PDF; not exact Word pagination.','word_export':'Word COM fixed-format export did not complete and held the generated document read-only. Only the known review-created Word process (28232) and its export PowerShell were stopped. No Word PDF was represented as successful.','rendering_attempts':'First regeneration met the Word file lock; corrected by stopping that review-owned instance. First ReportLab layout rejected a tall screenshot; image dimensions were bounded to the page frame and the complete render passed.','optional_renderer_download':'pymupdf 63 MB download stalled and was stopped; pypdfium2 used for visual inspection instead.'}
(base/'evidence/pdf-export.json').write_text(json.dumps(record,indent=2),encoding='utf8')
print(json.dumps(record,indent=2))
